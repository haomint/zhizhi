from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / 'work' / 'project_proposal_content.json').read_text(encoding='utf-8'))
OUT = ROOT / 'outputs'
OUT.mkdir(exist_ok=True)
DOCX = OUT / '知枝私人学习助手项目方案.docx'
MARKDOWN = OUT / '知枝私人学习助手项目方案.md'


def font(run, name='宋体', size=11, bold=False, color='000000'):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    props = run._element.get_or_add_rPr()
    fonts = props.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        props.insert(0, fonts)
    for key in ('eastAsia', 'ascii', 'hAnsi', 'cs'):
        fonts.set(qn('w:' + key), name)


def style_font(style, name, size, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    props = style.element.get_or_add_rPr()
    fonts = props.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        props.insert(0, fonts)
    for attr in ('eastAsia', 'ascii', 'hAnsi'):
        fonts.set(qn('w:' + attr), name)
    ppr = style.element.find(qn('w:pPr'))
    if ppr is not None:
        borders = ppr.find(qn('w:pBdr'))
        if borders is not None:
            ppr.remove(borders)


doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin, section.bottom_margin = Mm(17), Mm(17)
section.left_margin, section.right_margin = Mm(20), Mm(20)
section.footer_distance = Mm(7)
style_font(doc.styles['Normal'], '宋体', 11)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.38
normal.space_after = Pt(6)
normal.widow_control = True
style_font(doc.styles['Title'], '微软雅黑', 21, True)
doc.styles['Title'].paragraph_format.space_after = Pt(9)
doc.styles['Title'].paragraph_format.line_spacing = 1.1
style_font(doc.styles['Subtitle'], '微软雅黑', 12)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(10)
for name, size in [('Heading 1', 16), ('Heading 2', 11.5)]:
    style_font(doc.styles[name], '微软雅黑', size, True)
    fmt = doc.styles[name].paragraph_format
    fmt.space_before = Pt(9 if name == 'Heading 2' else 0)
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.2
    fmt.keep_with_next = True


def para(text, lead=False, flow=False):
    p = doc.add_paragraph()
    if flow:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.line_spacing = 1.6
        font(p.add_run(text), '微软雅黑', 10, color='344C42')
    elif lead:
        p.paragraph_format.space_before = Pt(4)
        font(p.add_run(text), '宋体', 11, bold=True)
    else:
        font(p.add_run(text))
    return p


def add_table(block):
    table = doc.add_table(rows=1, cols=len(block['headers']))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, width in zip(table.columns, block['widths']):
        col.width = Mm(width)
    props = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side)
        for name, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            e.set(qn('w:' + name), value)
        borders.append(e)
    props.append(borders)
    for ri, values in enumerate([block['headers']] + block['rows']):
        row = table.rows[0] if ri == 0 else table.add_row()
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        trpr.append(no_split)
        if ri == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for ci, (cell, value, width) in enumerate(zip(row.cells, values, block['widths'])):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, amount in [('top', 84), ('bottom', 84), ('left', 110), ('right', 110)]:
                margin = OxmlElement('w:' + side)
                margin.set(qn('w:w'), str(amount))
                margin.set(qn('w:type'), 'dxa')
                margins.append(margin)
            tcpr.append(margins)
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), '263E50' if ri == 0 else ('F3F5F7' if ri % 2 == 0 else 'FFFFFF'))
            tcpr.append(shading)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.18
            p.paragraph_format.keep_with_next = False
            if ri == 0 or block['headers'][ci] == '当前状态':
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            size = 9.4 if len(block['rows']) >= 10 else 10
            font(p.add_run(value), '微软雅黑' if ri == 0 else '宋体', size, bold=ri == 0,
                 color='FFFFFF' if ri == 0 else '000000')
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 0.4
    p.add_run().font.size = Pt(2)


footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run('第 '), '宋体', 9, color='666666')
for instruction, trailing in [('PAGE', ' 页  共 '), ('NUMPAGES', ' 页')]:
    r = footer.add_run()
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), instruction)
    r._r.addnext(field)
    font(footer.add_run(trailing), '宋体', 9, color='666666')

doc.add_paragraph(DATA['title'], style='Title')
doc.add_paragraph(DATA['subtitle'], style='Subtitle')
p = doc.add_paragraph()
font(p.add_run(DATA['date'] + '    ' + DATA['version']), '宋体', 9, color='666666')
p.paragraph_format.space_after = Pt(12)

for index, page in enumerate(DATA['pages']):
    heading = doc.add_paragraph(page['title'], style='Heading 1')
    heading.paragraph_format.page_break_before = index > 0
    for block in page['blocks']:
        kind = block['type']
        if kind == 'table':
            add_table(block)
        elif kind == 'h2':
            doc.add_paragraph(block['text'], style='Heading 2')
        elif kind == 'bullets':
            for item in block['items']:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Mm(4)
                p.paragraph_format.first_line_indent = Mm(-4)
                p.paragraph_format.space_after = Pt(4)
                font(p.add_run('• ' + item))
        else:
            para(block['text'], lead=kind == 'lead', flow=kind == 'flow')

doc.core_properties.title = DATA['title']
doc.core_properties.subject = DATA['subtitle']
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.core_properties.comments = ''
update = OxmlElement('w:updateFields')
update.set(qn('w:val'), 'true')
doc.settings.element.append(update)
doc.save(DOCX)

md = ['# ' + DATA['title'], '', DATA['subtitle'], '', DATA['date'] + '  ' + DATA['version'], '']
for page in DATA['pages']:
    md += ['## ' + page['title'], '']
    for block in page['blocks']:
        kind = block['type']
        if kind == 'h2':
            md += ['### ' + block['text'], '']
        elif kind == 'table':
            md += ['| ' + ' | '.join(block['headers']) + ' |', '| ' + ' | '.join('---' for _ in block['headers']) + ' |']
            md += ['| ' + ' | '.join(row) + ' |' for row in block['rows']]
            md += ['']
        elif kind == 'bullets':
            md += ['- ' + x for x in block['items']] + ['']
        elif kind == 'lead':
            md += ['**' + block['text'] + '**', '']
        else:
            md += [block['text'], '']
MARKDOWN.write_text('\n'.join(md), encoding='utf-8')
print(DOCX)
print(MARKDOWN)
print('Sections:', len(DATA['pages']), 'Tables:', len(doc.tables))
